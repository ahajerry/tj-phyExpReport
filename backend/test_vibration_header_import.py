"""Generate real XLSX files with independently sized tables and shifted headers."""
from io import BytesIO
from decimal import Decimal
import json
from pathlib import Path
import unittest

from openpyxl import Workbook, load_workbook

from .editable_instrument_data import editable_import
from .experiment_registry import validate_report
from .table_uploads import parse_table
from .test_table_uploads import resonance_export


CASES = [(2, 1, 3, 1), (7, 2, 17, 4), (31, 4, 41, 2), (103, 11, 121, 3)]


def variable_export(free_rows, trials, forced_rows, padding=1):
    workbook = Workbook()
    ws = workbook.active
    ws.title = "学生实验数据"
    row = padding

    def write(values):
        nonlocal row
        for col, value in values.items():
            ws.cell(row, col + padding, value)
        row += 1

    write({1: "波尔共振实验报告"})
    write({1: "姓名:", 2: "动态测试", 5: "学号:", 6: "1234567"})
    write({1: "表 1 自由震荡时的振幅 θ 与周期 T"})
    write({1: "振幅θ", 2: "周期T(s)", 4: "振幅θ", 5: "周期T(s)", 7: "振幅θ", 8: "周期T(s)"})
    for i in range(free_rows):
        write({col: value for pair, col0 in enumerate((1, 4, 7))
               for col, value in ((col0, 160 - (pair * free_rows + i) * .1), (col0 + 1, "1.5700"))})
        if i == 1:
            row += padding  # Interior blank rows must not truncate the table.
    write({1: "表 2 阻尼系数 β 的计算"})
    write({8: "阻尼档位", 9: 1})
    for trial in range(trials):
        write({1: "次数", 2: "振幅θₙ(°)", 4: "次数", 5: "振幅θₙ(°)", 7: "ln(θₙ/θₙ₊₅)"})
        for i in range(5):
            write({1: i + 1, 2: 150 - i * 10, 4: i + 6, 5: (150 - i * 10) / 2 + i + trial})
            row += padding if i == 2 else 0
        write({1: "ln(θₙ/θₙ₊₅)平均值", 7: .6931})
        write({3: "10T =", 4: "15.7000", 5: "秒", 6: "T =", 7: 1.57, 8: "秒"})
    write({1: "表 3 幅频特性和相频特性测量"})
    # Different order and spacing from both supplied templates.
    write({1: "相位差φ(°)\n读取值", 3: "振幅θ(°)\n测量值", 6: "强迫力矩周期(s)"})
    for i in range(forced_rows):
        write({1: 0 if i == 0 else 90, 3: 110 - i * .1, 6: "1.5800"})
        row += padding if i == 1 else 0
    write({1: "利用表 3 记录的数据，将计算结果填入表 4"})
    write({1: "表 4"})
    write({1: "强迫力矩周期(s)", 3: "相位差φ(°)读取值", 5: "θ(°)测量值", 6: "ω/ωr", 8: "相位差φ(°)计算值"})
    # Deliberately a different length and values; calculated tables must not
    # overwrite or duplicate raw measurements.
    for _ in range(forced_rows + 2):
        write({1: "9.9999", 3: 180, 5: 1, 6: 1, 8: 99})
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


class HeaderImportTests(unittest.TestCase):
    def test_ten_cycle_header_converts_once_for_editing_and_calculation(self):
        workbook = load_workbook(BytesIO(variable_export(3, 1, 5)))
        sheet = workbook.active
        for row in sheet:
            for cell in row:
                if cell.value == '强迫力矩周期(s)':
                    cell.value = '强迫力矩周期 10T(s)'
                    for index in range(cell.row + 1, sheet.max_row + 1):
                        if sheet.cell(index, cell.column).value == '1.5800':
                            sheet.cell(index, cell.column, '15.8000')
                    break
            else:
                continue
            break
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        parsed = parse_table('forced-vibration', 'ten-period.xlsx', output.getvalue())
        self.assertEqual(parsed['data']['drive_period'], ['1.58'] * 5)
        editable = editable_import({'experimentId': 'forced-vibration', **parsed})
        self.assertEqual(editable['data']['drive_time_10'], ['15.80'] * 5)
        self.assertNotIn('drive_period', editable['data'])
        self.assertEqual(validate_report({'experimentId': 'forced-vibration', **parsed}),
                         validate_report({'experimentId': 'forced-vibration', **editable}))

    def check_export(self, content, free_rows, trials, forced_rows):
        parsed = parse_table("forced-vibration", "variable.xlsx", content)
        self.assertEqual({key: len(values) for key, values in parsed["data"].items()},
                         dict(free_angle=free_rows * 3, free_period=free_rows * 3,
                              damped_peak=trials * 10, damped_period=trials,
                              drive_period=forced_rows, amplitude=forced_rows, phase=forced_rows))
        self.assertEqual(Decimal(parsed["data"]["free_angle"][-1]), Decimal(str(160 - (free_rows * 3 - 1) * .1)))
        self.assertEqual(parsed["data"]["drive_period"], ["1.5800"] * forced_rows)
        self.assertEqual(parsed["data"]["damped_period"], ["1.57"] * trials)
        self.assertEqual(parsed["data"]["phase"][0], "0")
        self.assertEqual(parsed["reportInfo"]["student"], "动态测试")
        self.assertTrue(validate_report({"experimentId": "forced-vibration", **parsed}))
        editable = editable_import({"experimentId": "forced-vibration", **parsed})
        self.assertTrue(validate_report({"experimentId": "forced-vibration", **editable}))
        self.assertEqual([len(table["rows"]) for table in parsed["sourceTables"]],
                         [free_rows, trials * 5, forced_rows])

    def test_independently_sized_tables_and_shifted_columns(self):
        for free, trials, forced, padding in CASES:
            with self.subTest(free=free, trials=trials, forced=forced, padding=padding):
                self.check_export(variable_export(free, trials, forced, padding), free, trials, forced)

    def test_instrument_layout_moves_beyond_old_search_window(self):
        workbook = load_workbook(BytesIO(resonance_export()))
        ws = workbook.active
        ws.insert_rows(39, 135)
        ws.insert_rows(1, 10)
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        parsed = parse_table("forced-vibration", "shifted.xlsx", output.getvalue())
        self.assertEqual(len(parsed["data"]["free_period"]), 90)
        self.assertEqual(parsed["data"]["damped_period"], ["1.5", "1.6", "1.7"])
        self.assertEqual(len(parsed["data"]["drive_period"]), 20)

    def test_incomplete_dynamic_rows_are_rejected(self):
        for header, column in (("周期T(s)", 2), ("强迫力矩周期(s)", 6)):
            workbook = load_workbook(BytesIO(variable_export(7, 2, 17)))
            ws = workbook.active
            target = next(cell for row in ws for cell in row if cell.value == header)
            ws.cell(target.row + 1, column + 1).value = None
            output = BytesIO()
            workbook.save(output)
            workbook.close()
            with self.subTest(header=header), self.assertRaises(ValueError):
                parse_table("forced-vibration", "partial.xlsx", output.getvalue())

    def test_only_student_data_is_read_when_blank_template_is_active(self):
        workbook = load_workbook(BytesIO(variable_export(7, 2, 17)))
        template = workbook.create_sheet("实验报告模板", 0)
        workbook.active = template
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        self.check_export(output.getvalue(), 7, 2, 17)

    def test_missing_student_sheet_never_falls_back_to_filled_template(self):
        workbook = load_workbook(BytesIO(variable_export(7, 2, 17)))
        workbook.active.title = "实验报告模板"
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        with self.assertRaisesRegex(ValueError, "缺少.*学生实验数据"):
            parse_table("forced-vibration", "template-only.xlsx", output.getvalue())

    def test_blank_student_sheet_never_reads_filled_template(self):
        workbook = load_workbook(BytesIO(variable_export(7, 2, 17)))
        workbook.active.title = "实验报告模板"
        workbook.create_sheet("学生实验数据")
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        with self.assertRaises(ValueError):
            parse_table("forced-vibration", "blank-student.xlsx", output.getvalue())


if __name__ == "__main__":
    import sys
    if len(sys.argv) == 3 and sys.argv[1] == "--generate":
        directory = Path(sys.argv[2])
        directory.mkdir(parents=True, exist_ok=True)
        summary = []
        checker = HeaderImportTests()
        for free, trials, forced, padding in CASES:
            name = f"free{free}_damping{trials}_forced{forced}.xlsx"
            path = directory / name
            path.write_bytes(variable_export(free, trials, forced, padding))
            checker.check_export(path.read_bytes(), free, trials, forced)
            summary.append(dict(file=name, free=free * 3, damping=trials * 10, forced=forced, result="PASS"))
        (directory / "results.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps(summary, ensure_ascii=False, indent=2))
    else:
        unittest.main()
