"""Instrument file imports use the actual export layouts supplied for this project."""

import json
import math
import re
import tempfile
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
import unittest
from unittest.mock import Mock
from openpyxl import Workbook

from .main import ReportHandler
from .millikan_exp import report_fields as millikan_fields
from .moment_of_inertia_exp import report_fields as torsion_fields
from .table_uploads import parse_table
from .experiment_registry import validate_report
from .forced_vibration_exp import uploaded_template, editable_source_tables
from .fill_sound_speed import fill_report, FIELD_PATTERN, COVER_PATTERN, PLOT_PATTERN
from scripts.verify_reports import samples


SOURCE_DIR = Path("D:/ProgramData/QQDownloads")


def sample(pattern):
    return next(SOURCE_DIR.glob(pattern), None)


def resonance_export(changes=None):
    """In-memory instrument layout, independent of the user's private workbook."""
    workbook = Workbook()
    ws = workbook.active
    ws.title = "学生实验数据"
    for cell, value in {"A1": "波尔共振实验数据", "A6": "1 振幅与周期关系",
                        "A40": "2 阻尼系数β的测量", "G41": "阻尼档位", "H41": 1,
                        "C43": "10T", "D62": "强迫力矩周期(s)", "H62": "相位差φ",
                        "B4": "测试学生", "F3": "1234567", "F4": "2024-08-06 17:12:51"}.items():
        ws[cell] = value
    for row in range(9, 39):
        for col in "ADG":
            ws[f"{col}{row}"] = 160 - row
        for col in "BEH":
            ws[f"{col}{row}"] = "1.5000"
    for trial, start in enumerate((44, 49, 54)):
        ws[f"C{start}"] = str(15 + trial)
        for offset in range(5):
            ws[f"D{start + offset}"] = offset + 1
            ws[f"F{start + offset}"] = 100 - offset * 10
            ws[f"H{start + offset}"] = (100 - offset * 10) / 2
    for row in range(63, 83):
        ws[f"A{row}"] = row - 62
        ws[f"D{row}"] = "1.5500"
        ws[f"F{row}"] = row
        ws[f"H{row}"] = 0 if row == 63 else 100
    for cell, value in (changes or {}).items():
        ws[cell] = value
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


class TableUploadTests(unittest.TestCase):
    def test_manual_submission_without_table(self):
        for experiment_id in ("millikan", "moment-of-inertia", "forced-vibration"):
            payload = {"experimentId": experiment_id, **samples[experiment_id]}
            encoded = json.dumps(payload, ensure_ascii=False).encode("utf-8")
            boundary = "manual-readings"
            multipart = (f'--{boundary}\r\nContent-Disposition: form-data; name="payload"\r\n\r\n'.encode()
                         + encoded + f"\r\n--{boundary}--\r\n".encode())
            for content_type, body in (("application/json", encoded),
                                       (f"multipart/form-data; boundary={boundary}", multipart)):
                with self.subTest(experiment_id=experiment_id, content_type=content_type):
                    fake = SimpleNamespace(headers={"Content-Type": content_type,
                                                    "Content-Length": str(len(body)), "X-Report-Request": "1"},
                                           rfile=BytesIO(body))
                    result, images = ReportHandler._read_submission(fake)
                    self.assertEqual(result, payload)
                    self.assertEqual(images, [])
                    self.assertTrue(validate_report(result))

    def test_csv_accepts_eight_and_rejects_nine_or_ten_drops(self):
        for count in (3, 8, 9, 10):
            content = ("当前电压,平衡电压,时间\n" + "120,120,10\n" * (count * 4)).encode()
            if count > 8:
                with self.assertRaisesRegex(ValueError, "3 至 8"):
                    parse_table("millikan", "data.csv", content)
            else:
                parsed = parse_table("millikan", "data.csv", content)
                self.assertTrue(all(len(values) == 8 for values in parsed["data"].values()))
                self.assertEqual(parsed["data"]["voltage"][count:], [""] * (8 - count))
                self.assertEqual(len(parsed["sourceTables"][0]["rows"]), count * 4)

    def test_json_requires_eight_slots(self):
        for count in (9, 10):
            payload = {"experimentId": "millikan", "parameters": {}, "data": {
                key: ["120" if key == "voltage" else "10"] * count
                for key in ("voltage", "time_1", "time_2", "time_3", "time_4")}}
            with self.assertRaisesRegex(ValueError, "8 个读数"):
                millikan_fields(payload)

    def test_millikan_csv_sample_has_eight_complete_drops(self):
        path = sample("密里根油滴.csv")
        if path is None:
            self.skipTest("user CSV is unavailable")
        parsed = parse_table("millikan", path.name, path.read_bytes())
        self.assertEqual(parsed["data"]["voltage"][0], "122.925")
        self.assertEqual(len(parsed["data"]["voltage"]), 8)
        self.assertEqual(len(parsed["sourceTables"][0]["rows"]), 32)
        fields = millikan_fields({"experimentId": "millikan", **parsed})
        self.assertIn("electron_charge_ls", fields)
        self.assertNotIn("charge_09", fields)

    def test_torsion_xlsx_sample_converts_units(self):
        path = sample("NB00000002*.xlsx")
        if path is None:
            self.skipTest("user XLSX is unavailable")
        parsed = parse_table("moment-of-inertia", path.name, path.read_bytes())
        self.assertEqual(parsed["parameters"]["mass_cylinder_g"], "717.3")
        self.assertEqual(parsed["parameters"]["slider_self_inertia"], "0.000040695")
        self.assertEqual(parsed["data"]["time_disk"], ["12.79", "12.79", "12.81"])
        fields = torsion_fields({"experimentId": "moment-of-inertia", **parsed})
        self.assertEqual(fields["slider_pair_inertia"], "0.000081390")

    def test_rejects_wrong_file_type_and_missing_csv_groups(self):
        with self.assertRaisesRegex(ValueError, "CSV"):
            parse_table("millikan", "data.xlsx", b"test")
        with self.assertRaisesRegex(ValueError, "连续 4 行"):
            parse_table("millikan", "data.csv", "平衡电压,时间\n1,2,3\n".encode())

    def test_resonance_uses_all_readings_and_three_individual_timings(self):
        parsed = parse_table("forced-vibration", "data.XLSX", resonance_export())
        self.assertEqual(len(parsed["data"]["free_period"]), 90)
        self.assertEqual(parsed["data"]["damped_period"], ["1.5", "1.6", "1.7"])
        self.assertEqual(len(parsed["data"]["damped_peak"]), 30)
        self.assertEqual(parsed["data"]["phase"][0], "0")
        fields = validate_report({"experimentId": "forced-vibration", **parsed})
        expected = sum(math.log(2) / (5 * period) for period in (1.5, 1.6, 1.7)) / 3
        self.assertAlmostEqual(float(fields["damping_beta_s_inv"]), expected, delta=0.0005)
        self.assertEqual(fields["free_period_90"], "1.5000")
        self.assertEqual(float(fields["peak_amplitude_deg"]), 82)
        self.assertEqual(len(re.findall(r"\([^()]+\)", fields["amplitude_graph_points"])), 20)
        self.assertEqual([len(table["rows"]) for table in parsed["sourceTables"]], [30, 15, 20])
        self.assertEqual(parsed["reportInfo"]["date"], "2024-08-06")

    def test_resonance_rejects_wrong_layout_missing_values_and_invalid_numbers(self):
        for changes in ({"A1": "扭摆"}, {"C43": "T"}, {"B9": None},
                        {"C49": 0}, {"H82": 361}, {"H82": "NaN"}, {"A82": 19}):
            with self.subTest(changes=changes), self.assertRaises(ValueError):
                parse_table("forced-vibration", "data.xlsx", resonance_export(changes))
        with self.assertRaisesRegex(ValueError, "损坏"):
            parse_table("forced-vibration", "data.xlsx", b"broken")

    def test_resonance_expanded_template_and_source_records_fill_completely(self):
        parsed = parse_table("forced-vibration", "data.xlsx", resonance_export())
        fields = validate_report({"experimentId": "forced-vibration", **parsed})
        source = uploaded_template()
        allowed = set(FIELD_PATTERN.findall(source)) | set(COVER_PATTERN.findall(source)) | set(PLOT_PATTERN.findall(source))
        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parents[1]) as directory:
            template = Path(directory) / "template.tex"
            template.write_text(source, encoding="utf-8")
            filled = fill_report({key: value for key, value in fields.items() if key in allowed},
                                 Path(directory) / "report.tex", template,
                                 source_tables=parsed["sourceTables"]).read_text(encoding="utf-8")
        for marker in ("free_period_90", "damped_peak_30", "damped_period_03", "table_drive_period_20", "table_frequency_ratio_20", "table_phase_calculated_20"):
            self.assertIn(r"\SetField{" + marker + "}", filled)
        self.assertEqual(filled.count(r"\begin{longtable}"), 6)
        self.assertNotIn(r"\PlotCoordinates", filled)

    def test_resonance_user_sample(self):
        path = sample("BEGZ000001*.xlsx")
        if path is None:
            self.skipTest("user XLSX is unavailable")
        parsed = parse_table("forced-vibration", path.name, path.read_bytes())
        self.assertEqual(parsed["data"]["free_angle"][30], "118")
        self.assertEqual(parsed["data"]["free_period"][-1], "1.5716")
        self.assertEqual(parsed["data"]["damped_period"], ["1.567", "1.5667", "1.5671"])
        self.assertEqual(parsed["data"]["damped_peak"][:10], ["142", "127", "114", "102", "90", "80", "72", "63", "56", "50"])
        self.assertEqual(parsed["data"]["drive_period"][-1], "1.5287")
        self.assertTrue(validate_report({"experimentId": "forced-vibration", **parsed}))

    def test_http_resonance_table_replaces_manual_readings_and_preserves_cover_override(self):
        boundary = "resonance-export"
        payload = {"experimentId": "forced-vibration", "data": {}, "parameters": {},
                   "reportInfo": {"student": "自填姓名", "student_id": "", "group": "2"}}
        body = (f'--{boundary}\r\nContent-Disposition: form-data; name="payload"\r\n\r\n'
                + json.dumps(payload, ensure_ascii=False) + "\r\n"
                + f'--{boundary}\r\nContent-Disposition: form-data; name="table"; filename="data.xlsx"\r\n'
                + "Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n").encode()
        body += resonance_export() + f"\r\n--{boundary}--\r\n".encode()
        fake = SimpleNamespace(headers={"Content-Type": f"multipart/form-data; boundary={boundary}",
                                        "Content-Length": str(len(body)), "X-Report-Request": "1"},
                               rfile=BytesIO(body))
        result, images = ReportHandler._read_submission(fake)
        self.assertEqual(images, [])
        self.assertEqual(result["reportInfo"]["student"], "自填姓名")
        self.assertEqual(result["reportInfo"]["student_id"], "1234567")
        self.assertEqual(result["reportInfo"]["group"], "2")
        self.assertEqual(len(result["data"]["drive_period"]), 20)
        self.assertTrue(validate_report(result))

    def test_resonance_import_returns_editable_values_without_creating_report(self):
        parsed = parse_table("forced-vibration", "data.xlsx", resonance_export())
        handler = Mock()
        handler.path = "/api/reports/import-table"
        handler._identity.return_value = ("visitor", False)
        handler._read_submission.return_value = ({"experimentId": "forced-vibration", **parsed}, [])
        ReportHandler.do_POST(handler)
        handler.queue.create.assert_not_called()
        status, result, _ = handler._json.call_args.args
        self.assertEqual(status, 200)
        self.assertEqual(result["data"]["damped_time_10"], ["15.0", "16.0", "17.0"])
        self.assertNotIn("damped_period", result["data"])
        result["data"]["amplitude"][-1] = "200"
        result["data"]["damped_time_10"][2] = "20.0"
        payload = {"experimentId": "forced-vibration", **result}
        fields = validate_report(payload)
        self.assertEqual(float(fields["peak_amplitude_deg"]), 200)
        expected = sum(math.log(2) / (5 * period) for period in (1.5, 1.6, 2)) / 3
        self.assertAlmostEqual(float(fields["damping_beta_s_inv"]), expected, delta=0.0005)
        records = editable_source_tables(payload)
        self.assertEqual(records[2]["rows"][-1][2], "200")
        self.assertEqual(records[1]["rows"][-1][1], "20.0")
        result["data"]["damped_time_10"][2] = "0"
        with self.assertRaisesRegex(ValueError, "正数"):
            validate_report(payload)

    def test_resonance_import_rejects_submission_without_uploaded_table(self):
        handler = Mock()
        handler.path = "/api/reports/import-table"
        handler._identity.return_value = ("visitor", False)
        handler._read_submission.return_value = ({"experimentId": "forced-vibration", **samples["forced-vibration"]}, [])
        ReportHandler.do_POST(handler)
        self.assertEqual(handler._json.call_args.args[0], 400)
        handler.queue.create.assert_not_called()

    def test_http_submission_uses_uploaded_table(self):
        path = sample("密里根油滴.csv")
        if path is None:
            self.skipTest("user CSV is unavailable")
        boundary = "instrument-export"
        payload = {"experimentId": "millikan", "reportInfo": {"student": "测试"}}
        body = (f"--{boundary}\r\nContent-Disposition: form-data; name=\"payload\"\r\n\r\n"
                + json.dumps(payload, ensure_ascii=False) + "\r\n"
                + f"--{boundary}\r\nContent-Disposition: form-data; name=\"table\"; filename=\"data.csv\"\r\n"
                + "Content-Type: text/csv\r\n\r\n").encode("utf-8") + path.read_bytes() + f"\r\n--{boundary}--\r\n".encode()
        fake = SimpleNamespace(headers={"Content-Type": f"multipart/form-data; boundary={boundary}",
                                        "Content-Length": str(len(body)), "X-Report-Request": "1"},
                               rfile=BytesIO(body))
        result, images = ReportHandler._read_submission(fake)
        self.assertEqual(images, [])
        self.assertEqual(result["reportInfo"]["student"], "测试")
        self.assertEqual(result["data"]["voltage"][0], "122.925")


if __name__ == "__main__":
    unittest.main()
