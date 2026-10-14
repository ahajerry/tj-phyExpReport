"""Parse instrument exports into report readings and source tables."""

from __future__ import annotations

import csv
import re
import unicodedata
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
from statistics import mean
from zipfile import BadZipFile

from openpyxl import load_workbook


MAX_TABLE_BYTES = 2 * 1024 * 1024


def _number(value, label):
    try:
        result = Decimal(str(value).strip())
    except (TypeError, InvalidOperation) as exc:
        raise ValueError(f"{label} 缺少有效数字") from exc
    if not result.is_finite() or not 0 < result < Decimal("1e12"):
        raise ValueError(f"{label} 必须为正数")
    return str(value).strip()


def _reading(ws, cell):
    return _number(ws[cell].value, cell)


def _plain(value: Decimal) -> str:
    return format(value.normalize(), "f")


def _header(value):
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", str(value or ""))).lower()


def _vibration_headers(ws):
    """Locate columns by their labels, including repeated continuation columns."""
    headers = []
    for row in ws:
        labels = {cell.column: _header(cell.value) for cell in row if cell.value is not None}
        periods = [c for c, s in labels.items() if "周期" in s and "查表" not in s]
        phases = [c for c, s in labels.items() if "相位差" in s and "计算" not in s]
        angles = [c for c, s in labels.items()
                  if ("振幅" in s or "θ" in s)
                  and not any(word in s for word in ("ln", "查表", "计算", "/θ", "²"))]
        timings = [c for c, s in labels.items() if s.startswith("10t")]
        indices = [c for c, s in labels.items() if s in ("n", "次数", "序号")]
        if periods and phases and angles:
            headers.append((row[0].row, "forced", (periods[0], angles[0], phases[0], indices)))
        elif periods and angles and not set(periods).intersection(angles):
            pairs = []
            for index, col in enumerate(angles):
                next_angle = angles[index + 1] if index + 1 < len(angles) else ws.max_column + 1
                following = next((c for c in periods if col < c < next_angle), None)
                if following is None:
                    raise ValueError(f"第 {row[0].row} 行振幅与周期表头不成对")
                pairs.append((col, following))
            if pairs:
                headers.append((row[0].row, "free", pairs))
        elif len(angles) >= 2 and indices:
            headers.append((row[0].row, "damped", (angles[:2], indices, timings)))
    return headers


def _vibration_boundary(ws, start, header_rows):
    for row in range(start + 1, ws.max_row + 1):
        if row in header_rows or any(
            re.match(r"^(?:表\d+|\d+[.、]?(?:自由|振幅|阻尼|强迫|受迫|幅频))", _header(cell.value))
            and isinstance(cell.value, str) and not _header(cell.value).replace(".", "", 1).isdigit()
            for cell in ws[row]
        ):
            return row
    return ws.max_row + 1


def parse_table(experiment_id: str, filename: str, content: bytes) -> dict:
    if not content or len(content) > MAX_TABLE_BYTES:
        raise ValueError("数据表必须小于 2 MiB 且不能为空")
    suffix = filename.lower().rsplit(".", 1)[-1]
    if experiment_id == "millikan" and suffix == "csv":
        return _millikan(content)
    if experiment_id == "moment-of-inertia" and suffix == "xlsx":
        return _torsion(content)
    if experiment_id == "forced-vibration" and suffix == "xlsx":
        return _forced_vibration(content)
    raise ValueError("密立根油滴请上传 CSV，扭摆法和波尔共振请上传 XLSX")


def _forced_vibration(content: bytes) -> dict:
    try:
        workbook = load_workbook(BytesIO(content), data_only=True)
    except (KeyError, OSError, BadZipFile, ValueError) as exc:
        raise ValueError("XLSX 文件损坏或格式不正确") from exc
    try:
        if "学生实验数据" not in workbook.sheetnames:
            raise ValueError("XLSX 缺少“学生实验数据”工作表，请在该工作表填写测量数据")
        ws = workbook["学生实验数据"]
        if not any("波尔共振" in str(cell.value) for row in ws for cell in row):
            raise ValueError("XLSX 不是所给格式的波尔共振实验数据表")
        headers = _vibration_headers(ws)
        if not all(any(kind == required for _, kind, _ in headers) for required in ("free", "damped", "forced")):
            raise ValueError("XLSX 缺少自由振动、阻尼或受迫振动的测量表头")
        header_rows = {row for row, _, _ in headers}
        data = {key: [] for key in ("free_angle", "free_period", "damped_peak",
                                    "damped_period", "drive_period", "amplitude", "phase")}
        free_rows, damped_rows, forced_rows = [], [], []
        free_width = 0

        def reading(row, col):
            return _number(ws.cell(row, col).value, ws.cell(row, col).coordinate)

        def occupied(row, columns):
            return any(ws.cell(row, col).value not in (None, "") for col in columns)

        def adjacent_value(cell):
            # Labels may be merged across several columns.
            col = cell.column + 1
            for merged in ws.merged_cells.ranges:
                if cell.coordinate in merged:
                    col = merged.max_col + 1
                    break
            return ws.cell(cell.row, col)

        setting = None
        report_info = {}
        cover_labels = {"姓名": "student", "学号": "student_id", "班号": "group", "时间": "date", "日期": "date"}
        for row in ws:
            for cell in row:
                label = _header(cell.value).rstrip(":")
                if label == "阻尼档位" and setting is None:
                    target = adjacent_value(cell)
                    if target.value not in (None, ""):
                        setting = _number(target.value, target.coordinate)
                if label in cover_labels:
                    report_info[cover_labels[label]] = str(adjacent_value(cell).value or "")
        if setting is None:
            raise ValueError("XLSX 缺少阻尼档位读数")
        if "date" in report_info:
            report_info["date"] = report_info["date"][:10]

        # Table 4 repeats measured columns beside calculated results. Prefer raw
        # forced-vibration tables whenever present, regardless of their row count.
        raw_forced = {row for row, kind, _ in headers if kind == "forced"
                      and not any("计算" in _header(c.value) or "ω" in _header(c.value) for c in ws[row])}
        for start, kind, columns in headers:
            end = _vibration_boundary(ws, start, header_rows)
            if kind == "free":
                rows = []
                free_width = max(free_width, len(columns) * 2)
                for row in range(start + 1, end):
                    values = []
                    for first, second in columns:
                        values.extend([reading(row, first), reading(row, second)]
                                      if occupied(row, (first, second)) else ["", ""])
                    if any(values):
                        rows.append(values)
                free_rows.extend(rows)
                for offset in range(0, len(columns) * 2, 2):
                    data["free_angle"].extend(row[offset] for row in rows if row[offset])
                    data["free_period"].extend(row[offset + 1] for row in rows if row[offset])
            elif kind == "damped":
                angles, indices, timings = columns
                groups, group = [], []
                for row in range(start + 1, end):
                    if not occupied(row, angles):
                        continue
                    # Summary and timing rows below the peaks are not readings.
                    if any("平均" in str(cell.value) or _header(cell.value).startswith("10t")
                           for cell in ws[row]):
                        continue
                    index = Decimal(reading(row, indices[0]))
                    if index == 1 and group:
                        groups.append(group)
                        group = []
                    if index != len(group) + 1:
                        raise ValueError(f"阻尼峰值序号 {ws.cell(row, indices[0]).coordinate} 不连续")
                    if len(indices) > 1 and Decimal(reading(row, indices[1])) != index + 5:
                        raise ValueError("阻尼峰值第二列次数应为 n+5")
                    group.append((row, reading(row, angles[0]), reading(row, angles[1])))
                if group:
                    groups.append(group)
                # Some templates place the 10T timing below the peak table.
                below_timings = [adjacent_value(cell) for row in range(start + 1, end)
                                 for cell in ws[row] if _header(cell.value).startswith("10t")]
                for index, group in enumerate(groups):
                    if len(group) != 5:
                        raise ValueError("每次阻尼测量需要 n=1 至 5 与 n+5 的十个峰值")
                    if timings:
                        timing = reading(group[0][0], timings[0])
                    elif index < len(below_timings):
                        cell = below_timings[index]
                        timing = _number(cell.value, cell.coordinate)
                    else:
                        raise ValueError("XLSX 阻尼测量缺少 10T 计时表头或读数")
                    data["damped_period"].append(_plain(Decimal(timing) / 10))
                    data["damped_peak"].extend([item[1] for item in group] + [item[2] for item in group])
                    trial = str(len(data["damped_period"]))
                    damped_rows.extend([[trial, timing, str(i), first, second]
                                        for i, (_, first, second) in enumerate(group, 1)])
            elif not raw_forced or start in raw_forced:
                period_col, amplitude_col, phase_col, indices = columns
                local_count = 0
                for row in range(start + 1, end):
                    if not occupied(row, (period_col, amplitude_col, phase_col)):
                        continue
                    # Instruction lines between tables mark the end of records.
                    if any("利用表" in str(c.value) for c in ws[row]):
                        break
                    local_count += 1
                    if indices and Decimal(reading(row, indices[0])) != local_count:
                        raise ValueError(f"受迫振动序号 {ws.cell(row, indices[0]).coordinate} 不连续")
                    period, amplitude = reading(row, period_col), reading(row, amplitude_col)
                    if "10t" in _header(ws.cell(start, period_col).value):
                        period = _plain(Decimal(period) / 10)
                    cell = ws.cell(row, phase_col)
                    try:
                        phase = Decimal(str(cell.value).strip())
                    except InvalidOperation as exc:
                        raise ValueError(f"{cell.coordinate} 缺少有效相位差") from exc
                    if not phase.is_finite() or not 0 <= phase <= 180:
                        raise ValueError(f"{cell.coordinate} 相位滞后角读数应在 0 至 180 度之间")
                    phase_text = str(cell.value).strip()
                    for key, value in (("drive_period", period), ("amplitude", amplitude), ("phase", phase_text)):
                        data[key].append(value)
                    forced_rows.append([str(len(forced_rows) + 1), period, amplitude, phase_text])
        if any(not values for values in data.values()):
            raise ValueError("XLSX 测量表缺少有效数据行")
        free_rows = [row + [""] * (free_width - len(row)) for row in free_rows]
        return {"data": data, "parameters": {"damping_setting": setting}, "reportInfo": report_info,
                "sourceTables": [
                    {"title": "波尔共振 XLSX 自由振动原始记录（各列依次续接）",
                     "headers": [label for i in range(1, free_width // 2 + 1)
                                 for label in (f"振幅{i}/°", f"周期{i}/s")], "rows": free_rows},
                    {"title": "波尔共振 XLSX 阻尼原始记录（档位 " + setting + "）",
                     "headers": ["测量组", "10T/s", "n", "峰值n/°", "峰值(n+5)/°"], "rows": damped_rows},
                    {"title": "波尔共振 XLSX 受迫振动原始记录",
                     "headers": ["序号", "驱动单周期/s", "稳定振幅/°", "滞后角正读数/°"], "rows": forced_rows}]}
    finally:
        workbook.close()


def _millikan(content: bytes) -> dict:
    try:
        source = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        source = content.decode("gb18030")
    rows = list(csv.reader(StringIO(source)))
    if not rows or "平衡电压" not in "".join(rows[0]) or "时间" not in "".join(rows[0]):
        raise ValueError("CSV 表头需包含平衡电压和时间")
    readings = [row for row in rows[1:] if any(cell.strip() for cell in row)]
    if len(readings) % 4 or not 3 <= len(readings) // 4 <= 8:
        raise ValueError("CSV 每颗油滴需有连续 4 行记录，共需 3 至 8 颗")
    data = {"voltage": [], **{f"time_{j}": [] for j in range(1, 5)}}
    source_rows = []
    for index in range(0, len(readings), 4):
        group = readings[index:index + 4]
        if any(len(row) < 3 for row in group):
            raise ValueError(f"第 {index + 2} 行缺少电压或时间")
        voltages = [_number(row[1], "平衡电压") for row in group]
        times = [_number(row[2], "下降时间") for row in group]
        data["voltage"].append(str(mean(Decimal(value) for value in voltages)))
        for j, value in enumerate(times, 1):
            data[f"time_{j}"].append(value)
        for j, row in enumerate(group, 1):
            source_rows.append([str(index // 4 + 1), str(j), row[0].strip(),
                                voltages[j - 1], times[j - 1], row[3].strip() if len(row) > 3 else ""])
    for values in data.values():
        values.extend([""] * (8 - len(values)))
    return {"data": data, "parameters": {}, "sourceTables": [{
        "title": "密立根油滴仪 CSV 原始记录",
        "headers": ["油滴", "次数", "当前电压/V", "平衡电压/V", "时间/s", "电荷量/C"], "rows": source_rows,
    }]}


def _torsion(content: bytes) -> dict:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        ws = workbook.active
        if "扭摆" not in str(ws["A1"].value) or "滑块质量" not in str(ws["A34"].value):
            raise ValueError("XLSX 不是所给格式的扭摆实验数据表")
        data = {}
        for key, row in (("cylinder", 16), ("disk", 17), ("tube", 29),
                         ("ball", 30), ("rod", 31)):
            data[f"time_{key}"] = [_reading(ws, f"{col}{row}") for col in "CDE"]
        for x, row in ((5, 37), (10, 38), (15, 39), (20, 40)):
            if float(_reading(ws, f"A{row}")) != x:
                raise ValueError(f"滑块距离 A{row} 与模板不一致")
            data[f"slider_time_{x}"] = [_reading(ws, f"{col}{row}") for col in "CDE"]
        def average_mm(row):
            return str((mean(Decimal(_reading(ws, f"{col}{row}")) for col in "CDE") * 10).quantize(Decimal("0.001")))
        parameters = {
            "mass_cylinder_g": _plain(Decimal(_reading(ws, "B12")) * 1000),
            "diameter_cylinder_mm": average_mm(12),
            "mass_tube_g": _plain(Decimal(_reading(ws, "B22")) * 1000),
            "diameter_tube_outer_mm": average_mm(23),
            "diameter_tube_inner_mm": average_mm(22),
            "mass_ball_g": _plain(Decimal(_reading(ws, "B24")) * 1000),
            "diameter_ball_mm": average_mm(24),
            "mass_rod_g": _plain(Decimal(_reading(ws, "B25")) * 1000),
            "length_rod_mm": average_mm(25),
            "slider_mass_g": _plain(mean(Decimal(_reading(ws, cell)) for cell in ("B34", "C34")) * 1000),
            "slider_self_inertia": str(Decimal(_reading(ws, "H34")) * Decimal("1e-4") / 2),
        }
        geometry = {key: [_reading(ws, f"{col}{row}") for col in "CDE"] for key, row in (
            ("cylinder", 12), ("tube_inner", 22), ("tube_outer", 23),
            ("ball", 24), ("rod", 25))}
        tables = [{"title": "扭摆仪 XLSX 质量与尺寸记录", "headers": ["物体", "质量/kg", "尺寸1/cm", "尺寸2/cm", "尺寸3/cm"],
                   "rows": [[str(ws[f"A{r}"].value or ""), str(ws[f"B{r}"].value or ""),
                             *[str(ws[f"{c}{r}"].value or "") for c in "CDE"]]
                            for r in (12, 22, 23, 24, 25)]},
                  {"title": "扭摆仪 XLSX 周期计时记录", "headers": ["物体/位置", "周期数", "时长1/s", "时长2/s", "时长3/s"],
                   "rows": [[str(ws[f"A{r}"].value or ("圆筒" if r == 29 else "")), str(ws[f"B{r}"].value or ""),
                             *[str(ws[f"{c}{r}"].value or "") for c in "CDE"]]
                            for r in (16, 17, 29, 30, 31, 37, 38, 39, 40)]},
                  {"title": "扭摆仪 XLSX 滑块参数", "headers": ["滑块1质量/kg", "滑块2质量/kg", "两滑块自身惯量/×10⁻⁴ kg·m²"],
                   "rows": [[str(ws["B34"].value), str(ws["C34"].value), str(ws["H34"].value)]]}]
        return {"data": data, "parameters": parameters, "geometryReadings": geometry,
                "sourceTables": tables,
                "reportInfo": {"student": str(ws["E3"].value or ""),
                               "student_id": str(ws["H3"].value or ""),
                               "date": str(ws["H4"].value or "")[:10]}}
    except (KeyError, OSError, BadZipFile, ValueError) as exc:
        if isinstance(exc, ValueError):
            raise
        raise ValueError("XLSX 文件损坏或格式不正确") from exc
